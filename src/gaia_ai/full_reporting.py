from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEET_NAMES = {
    "executive_summary": "Executive Summary",
    "liability_map": "Liability Map",
    "regime_sectors": "Regime and Sectors",
    "conviction": "GFSS and GCC",
    "narrative_agents": "Five Agents",
    "narrative_evidence": "Evidence Trace",
    "uncertainty_gate": "Uncertainty Gate",
    "allocation_research": "Research Allocation",
    "allocation_audit": "Allocation Audit",
    "risk_buckets": "Strategic Buckets",
    "risk_schedule": "Dynamic Risk Budget",
    "funding_scenarios": "Funding Scenarios",
    "ofc_thresholds": "OFC Thresholds",
    "payment_probability": "Payment Probability",
    "annual_paths": "Annual Percentiles",
    "commitments_2031": "2031 Commitments",
    "reserve_glidepath": "2033 Reserve",
    "gfss_sensitivity": "GFSS Sensitivity",
    "model_backtest": "Model Backtest",
    "model_coefficients": "Model Coefficients",
    "latest_predictions": "Latest Predictions",
}


def write_full_workbook(path: str | Path, frames: dict[str, pd.DataFrame], summary: dict) -> None:
    path = Path(path)
    readme = pd.DataFrame({
        "topic": [
            "System status", "What this is", "What this is not", "Human approval", "Narrative provider",
            "Evidence", "Return model", "Funding probability", "2031 commitment", "2033 reserve",
            "Trading", "Daily prices", "Reproducibility",
        ],
        "explanation": [
            str(summary["system_status"]),
            "Complete GAIA research workflow: client liabilities, regime, sector scores, quant ML, deep dive, five agents, uncertainty, GFSS/GCC, risk, allocation, and funding scenarios.",
            "Not an approved IPS, not a promise of returns, and not a fully automated investment committee.",
            "Uploaded 3D research is AI-simulated and human_approved=False. Unapproved candidates receive zero approved target weight.",
            "The extractive fallback is deterministic keyword/sentence extraction, not an LLM. Configure an OpenAI-compatible provider separately for a real LLM run.",
            "Every narrative score retains document ID, source URL, agent, dimension, confidence, direction, and an evidence field.",
            "Ridge regression is evaluated with expanding-window walk-forward tests. GFSS weights receive ±20% rank-sensitivity checks.",
            "OFC is the simulated probability that all ten payments are funded without sponsor support. It is assumption-dependent, not a guaranteed probability.",
            "The range is a distribution of contingent 2033 reserve shortfalls evaluated for communication in 2031; zero is a valid outcome.",
            "Reserve equals the present value of remaining beginning-of-year payments plus the flexibility buffer, funded before the first 2033 payment.",
            "Research output only. No broker login, no automatic order, and investment-committee approval remains mandatory.",
            "The system does not claim precise daily price forecasts.",
            "CSV and JSON outputs accompany this workbook; random seeds and all model assumptions are stored in the project configuration.",
        ],
    })
    sources = pd.DataFrame({
        "source_name": [
            "Wharton 2026 GAIA integrated ten-module master", "2026 WGY Gao case study",
            "GAIA 3C/3D workbook", "GAIA quarterly signal workbook", "allocation_12_plus_etf workbook",
            "Project data CSV/JSON files",
        ],
        "source_location_or_url": [
            "User-uploaded strategy blueprint", "User-uploaded client case", "User-uploaded; AI-simulated research",
            "User-uploaded; converted to project CSV", "User-uploaded strategic weights", "data/ in the project",
        ],
        "usage_note": [
            "Architecture and governance", "Cash-flow timing and client commitments", "Deep-dive inputs; not human-approved",
            "3B/3C quarterly panel and account snapshot", "60% SPYM / 40% stock-sleeve base", "Editable point-in-time inputs and research assumptions",
        ],
    })
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="READ ME", index=False)
        for key, sheet_name in SHEET_NAMES.items():
            if key in frames:
                frames[key].to_excel(writer, sheet_name=sheet_name, index=False)
        sources.to_excel(writer, sheet_name="Sources", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17324D")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 32
            for index, column in enumerate(ws.columns, 1):
                values = [str(cell.value) if cell.value is not None else "" for cell in column[:200]]
                width = min(48, max(11, max(map(len, values), default=10) + 2))
                ws.column_dimensions[get_column_letter(index)].width = width
            if ws.title in {"READ ME", "Evidence Trace", "Sources"}:
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
            if ws.title == "GFSS and GCC":
                headers = {cell.value: cell.column for cell in ws[1]}
                if "gcc" in headers and ws.max_row > 2:
                    letter = get_column_letter(headers["gcc"])
                    ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}", DataBarRule(start_type="min", end_type="max", color="5B9BD5"))
            if ws.title == "Funding Scenarios":
                headers = {cell.value: cell.column for cell in ws[1]}
                if "operating_funding_confidence" in headers and ws.max_row > 2:
                    letter = get_column_letter(headers["operating_funding_confidence"])
                    ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47"))
