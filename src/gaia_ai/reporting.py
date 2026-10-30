from __future__ import annotations

import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def json_default(value):
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return None if np.isnan(value) else float(value)
    if isinstance(value, (pd.Timestamp,)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, Path):
        return str(value)
    raise TypeError(f"Not JSON serializable: {type(value).__name__}")


def write_json(path: str | Path, payload: dict) -> None:
    with Path(path).open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2, default=json_default)


def write_daily_workbook(
    path: str | Path,
    daily: pd.DataFrame,
    instructions: pd.DataFrame,
    summary: dict,
    predictions: pd.DataFrame,
) -> None:
    """Write a portable Excel report; CSV/JSON remain the machine-readable sources."""
    path = Path(path)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_frame = pd.DataFrame({"item": list(summary), "value": list(summary.values())})
        summary_frame.to_excel(writer, sheet_name="Daily Summary", index=False)
        instructions.to_excel(writer, sheet_name="Trade Review", index=False)
        daily.to_excel(writer, sheet_name="Position Detail", index=False)
        predictions.to_excel(writer, sheet_name="Model Predictions", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17324D")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 30
            for index, column in enumerate(ws.columns, 1):
                values = [str(cell.value) if cell.value is not None else "" for cell in column[:200]]
                width = min(42, max(11, max(map(len, values), default=10) + 2))
                ws.column_dimensions[get_column_letter(index)].width = width
